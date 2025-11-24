#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
마이그레이션 체크리스트 엑셀 템플릿 생성 스크립트

이 스크립트는 마이그레이션 대상 리소스 목록을 관리하기 위한
엑셀 템플릿 파일을 생성합니다.
"""

import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
except ImportError:
    print("Error: openpyxl is required. Install it with: pip install openpyxl")
    sys.exit(1)


def create_template(filename: str = "migration_checklist_template.xlsx"):
    """
    마이그레이션 체크리스트 엑셀 템플릿 생성
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Migration List"

    # 컬럼 헤더
    headers = [
        '리소스 타입',
        '리소스 이름',
        '소스 계정',
        '대상 계정',
        '마이그레이션 상태',
        '검증 일시',
        '비고'
    ]

    # 컬럼 너비 설정
    column_widths = [20, 50, 15, 15, 20, 20, 40]

    # 스타일 정의
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_alignment = Alignment(horizontal="center", vertical="center")

    border_thin = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # 헤더 작성
    for col_idx, (header, width) in enumerate(zip(headers, column_widths), start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = border_thin
        ws.column_dimensions[cell.column_letter].width = width

    # 샘플 데이터 추가
    sample_data = [
        ['Lambda', 'my-function-1', '111111111111', '222222222222', '', '', ''],
        ['Lambda', 'my-function-2', '111111111111', '222222222222', '', '', ''],
        ['StepFunction', 'my-state-machine', '111111111111', '222222222222', '', '', ''],
        ['EventBridgeRule', 'my-rule', '111111111111', '222222222222', '', '', ''],
        ['EventBridgeSchedule', 'default/my-schedule', '111111111111', '222222222222', '', '', '스케줄 그룹명/스케줄명 형식'],
        ['APIGateway', 'my-api', '111111111111', '222222222222', '', '', ''],
        ['SecretsManager', 'my-secret', '111111111111', '222222222222', '', '', ''],
    ]

    # 샘플 데이터 스타일
    data_alignment = Alignment(horizontal="left", vertical="center")
    sample_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    for row_idx, row_data in enumerate(sample_data, start=2):
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = data_alignment
            cell.border = border_thin
            if row_idx <= 8:  # 샘플 데이터에만 배경색 적용
                cell.fill = sample_fill

    # 빈 행 추가 (사용자가 입력할 공간)
    for row_idx in range(9, 20):
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx, value='')
            cell.border = border_thin
            cell.alignment = data_alignment

    # 행 높이 설정
    for row in ws.iter_rows(min_row=1, max_row=20):
        ws.row_dimensions[row[0].row].height = 20

    # Instructions 시트 추가
    ws_instructions = wb.create_sheet("사용 방법")

    instructions = [
        ["마이그레이션 체크리스트 사용 방법", ""],
        ["", ""],
        ["1. 기본 정보", ""],
        ["이 엑셀 파일은 AWS 리소스 마이그레이션 대상을 관리하고", ""],
        ["실제 이관 완료 여부를 검증하기 위한 체크리스트입니다.", ""],
        ["", ""],
        ["2. 컬럼 설명", ""],
        ["• 리소스 타입", "Lambda, StepFunction, EventBridgeRule, EventBridgeSchedule, APIGateway, SecretsManager"],
        ["• 리소스 이름", "AWS 리소스의 이름 (ARN 아님)"],
        ["• 소스 계정", "소스 AWS 계정 ID (12자리 숫자)"],
        ["• 대상 계정", "대상 AWS 계정 ID (12자리 숫자)"],
        ["• 마이그레이션 상태", "검증 스크립트가 자동으로 업데이트"],
        ["• 검증 일시", "검증 스크립트가 자동으로 업데이트"],
        ["• 비고", "검증 스크립트가 자동으로 업데이트 (또는 수동 메모)"],
        ["", ""],
        ["3. 리소스 타입별 이름 형식", ""],
        ["• Lambda", "함수 이름 (예: my-function)"],
        ["• StepFunction", "상태 머신 이름 (예: my-state-machine)"],
        ["• EventBridgeRule", "규칙 이름 (예: my-rule)"],
        ["• EventBridgeSchedule", "그룹명/스케줄명 (예: default/my-schedule) 또는 스케줄명만"],
        ["• APIGateway", "REST API 이름 (예: my-api)"],
        ["• SecretsManager", "시크릿 이름 (예: my-secret)"],
        ["", ""],
        ["4. 검증 스크립트 실행", ""],
        ["python3 verify_migration_from_excel.py", ""],
        ["", ""],
        ["환경변수 설정:", ""],
        ["export DST_PROFILE=dst", "# 대상 계정 AWS 프로파일"],
        ["export AWS_REGION=ap-northeast-2", "# AWS 리전"],
        ["export EXCEL_FILE_PATH=migration_checklist.xlsx", "# 입력 파일"],
        ["export OUTPUT_FILE_PATH=migration_verification_result.xlsx", "# 출력 파일"],
        ["", ""],
        ["5. 결과 확인", ""],
        ["검증이 완료되면 '마이그레이션 상태' 컬럼이 업데이트됩니다:", ""],
        ["• 이관 완료: 녹색 배경", ""],
        ["• 이관 미완료: 빨간색 배경", ""],
        ["", ""],
        ["6. 주의사항", ""],
        ["• 'Migration List' 시트 이름을 변경하지 마세요", ""],
        ["• 헤더 행(첫 번째 행)을 삭제하지 마세요", ""],
        ["• 리소스 타입은 정확하게 입력해야 합니다 (대소문자 구분)", ""],
    ]

    for row_idx, (col1, col2) in enumerate(instructions, start=1):
        cell1 = ws_instructions.cell(row=row_idx, column=1, value=col1)
        cell2 = ws_instructions.cell(row=row_idx, column=2, value=col2)

        # 제목 스타일
        if row_idx == 1:
            cell1.font = Font(bold=True, size=14, color="4472C4")
        # 섹션 제목 스타일
        elif col1.startswith(('1.', '2.', '3.', '4.', '5.', '6.')):
            cell1.font = Font(bold=True, size=12)
        # 일반 텍스트
        else:
            cell1.font = Font(size=10)
            cell2.font = Font(size=10)

    ws_instructions.column_dimensions['A'].width = 60
    ws_instructions.column_dimensions['B'].width = 80

    # 파일 저장
    wb.save(filename)
    print(f"✅ 템플릿 파일이 생성되었습니다: {filename}")
    print(f"")
    print(f"다음 단계:")
    print(f"1. {filename} 파일을 열어서 마이그레이션 대상 리소스를 입력하세요")
    print(f"2. 입력이 완료되면 파일을 'migration_checklist.xlsx'로 저장하세요")
    print(f"3. verify_migration_from_excel.py 스크립트를 실행하세요")


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description='마이그레이션 체크리스트 템플릿 생성')
    parser.add_argument(
        '-o', '--output',
        default='migration_checklist_template.xlsx',
        help='출력 파일명 (기본값: migration_checklist_template.xlsx)'
    )

    args = parser.parse_args()
    create_template(args.output)
