#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import os
import json
import boto3
import logging
import csv
from typing import Dict, List, Optional, Set
from botocore.exceptions import ClientError
from datetime import datetime

# openpyxl for Excel operations (optional, only needed for .xlsx files)
OPENPYXL_AVAILABLE = False
try:
    from openpyxl import load_workbook, Workbook
    from openpyxl.styles import PatternFill, Font
    OPENPYXL_AVAILABLE = True
except ImportError:
    pass

# ===================================================
# 로깅 설정
# ===================================================
LOG_LEVEL = os.getenv("LOG_LEVEL", "INFO").upper()
logging.basicConfig(
    level=getattr(logging, LOG_LEVEL, logging.INFO),
    format='%(asctime)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

# ===================================================
# 설정
# ===================================================
DST_PROFILE = os.getenv("DST_PROFILE", "dst")
REGION = os.getenv("AWS_REGION", "ap-northeast-2")
EXCEL_FILE_PATH = os.getenv("EXCEL_FILE_PATH", "migration_checklist.xlsx")
OUTPUT_FILE_PATH = os.getenv("OUTPUT_FILE_PATH", "migration_verification_result.xlsx")

# 엑셀 시트 이름
SHEET_NAME = os.getenv("SHEET_NAME", "Migration List")

# 엑셀 컬럼 인덱스 (0-based)
COL_RESOURCE_TYPE = 0  # A: 리소스 타입
COL_RESOURCE_NAME = 1  # B: 리소스 이름
COL_SOURCE_ACCOUNT = 2  # C: 소스 계정 ID
COL_DEST_ACCOUNT = 3  # D: 대상 계정 ID
COL_STATUS = 4  # E: 마이그레이션 상태
COL_VERIFIED_AT = 5  # F: 검증 일시
COL_NOTES = 6  # G: 비고

# 지원 리소스 타입
SUPPORTED_RESOURCE_TYPES = [
    'Lambda',
    'StepFunction',
    'EventBridgeRule',
    'EventBridgeSchedule',
    'APIGateway',
    'SecretsManager'
]

# ===================================================
# AWS 클라이언트
# ===================================================
session_dst = boto3.Session(profile_name=DST_PROFILE, region_name=REGION)
lambda_client_dst = session_dst.client('lambda')
sfn_client_dst = session_dst.client('stepfunctions')
events_client_dst = session_dst.client('events')
scheduler_client_dst = session_dst.client('scheduler')
apigateway_client_dst = session_dst.client('apigateway')
secretsmanager_client_dst = session_dst.client('secretsmanager')


# ===================================================
# AWS 리소스 조회 함수
# ===================================================

def check_lambda_exists(function_name: str) -> bool:
    """Lambda 함수가 존재하는지 확인"""
    try:
        lambda_client_dst.get_function(FunctionName=function_name)
        return True
    except ClientError as e:
        if e.response['Error']['Code'] == 'ResourceNotFoundException':
            return False
        logger.error(f"Error checking Lambda {function_name}: {e}")
        return False


def check_stepfunction_exists(state_machine_name: str) -> bool:
    """Step Functions 상태 머신이 존재하는지 확인"""
    try:
        # List all state machines and check if name exists
        paginator = sfn_client_dst.get_paginator('list_state_machines')
        for page in paginator.paginate():
            for sm in page['stateMachines']:
                if sm['name'] == state_machine_name:
                    return True
        return False
    except ClientError as e:
        logger.error(f"Error checking StepFunction {state_machine_name}: {e}")
        return False


def check_eventbridge_rule_exists(rule_name: str) -> bool:
    """
    EventBridge Rule이 존재하는지 확인

    Note: 비활성화(DISABLED) 상태의 규칙도 존재하는 것으로 간주합니다.
    이관 완료 여부는 규칙의 존재 여부로만 판단하며, 활성화 상태는 고려하지 않습니다.
    """
    try:
        # describe_rule은 ENABLED/DISABLED 상태 모두 반환
        events_client_dst.describe_rule(Name=rule_name)
        return True
    except ClientError as e:
        if e.response['Error']['Code'] == 'ResourceNotFoundException':
            return False
        logger.error(f"Error checking EventBridge Rule {rule_name}: {e}")
        return False


def check_eventbridge_schedule_exists(schedule_name: str, group_name: str = 'default') -> bool:
    """
    EventBridge Scheduler가 존재하는지 확인

    Note: 비활성화(DISABLED) 상태의 스케줄도 존재하는 것으로 간주합니다.
    이관 완료 여부는 스케줄의 존재 여부로만 판단하며, 활성화 상태는 고려하지 않습니다.
    """
    try:
        logger.debug(f"Checking EventBridge Schedule - Name: '{schedule_name}', GroupName: '{group_name}'")
        # get_schedule은 ENABLED/DISABLED 상태 모두 반환
        response = scheduler_client_dst.get_schedule(
            Name=schedule_name,
            GroupName=group_name
        )
        state = response.get('State', 'UNKNOWN')
        logger.debug(f"EventBridge Schedule found - Name: '{schedule_name}', State: {state}")
        return True
    except ClientError as e:
        error_code = e.response['Error']['Code']
        if error_code == 'ResourceNotFoundException':
            logger.debug(f"EventBridge Schedule not found - Name: '{schedule_name}', GroupName: '{group_name}'")
            return False
        logger.error(f"Error checking EventBridge Schedule {schedule_name} in group {group_name}: [{error_code}] {e}")
        return False


def check_apigateway_exists(api_name: str) -> bool:
    """API Gateway REST API가 존재하는지 확인"""
    try:
        paginator = apigateway_client_dst.get_paginator('get_rest_apis')
        for page in paginator.paginate():
            for api in page['items']:
                if api['name'] == api_name:
                    return True
        return False
    except ClientError as e:
        logger.error(f"Error checking API Gateway {api_name}: {e}")
        return False


def check_secretsmanager_exists(secret_name: str) -> bool:
    """Secrets Manager 시크릿이 존재하는지 확인"""
    try:
        secretsmanager_client_dst.describe_secret(SecretId=secret_name)
        return True
    except ClientError as e:
        if e.response['Error']['Code'] == 'ResourceNotFoundException':
            return False
        logger.error(f"Error checking Secrets Manager {secret_name}: {e}")
        return False


def verify_resource(resource_type: str, resource_name: str) -> tuple[bool, str]:
    """
    리소스가 대상 계정에 존재하는지 확인
    Returns: (exists: bool, notes: str)
    """
    resource_name = resource_name.strip()

    if not resource_name:
        return False, "리소스 이름이 비어있음"

    try:
        if resource_type == 'Lambda':
            exists = check_lambda_exists(resource_name)
            return exists, "존재함" if exists else "미존재"

        elif resource_type == 'StepFunction':
            exists = check_stepfunction_exists(resource_name)
            return exists, "존재함" if exists else "미존재"

        elif resource_type == 'EventBridgeRule':
            exists = check_eventbridge_rule_exists(resource_name)
            return exists, "존재함" if exists else "미존재"

        elif resource_type == 'EventBridgeSchedule':
            # 스케줄 이름만 있거나, "그룹명/스케줄명" 형태로 제공
            if '/' in resource_name:
                group_name, schedule_name = resource_name.split('/', 1)
            else:
                group_name = 'default'
                schedule_name = resource_name

            logger.debug(f"EventBridgeSchedule verification - Input: '{resource_name}' -> Group: '{group_name}', Schedule: '{schedule_name}'")
            exists = check_eventbridge_schedule_exists(schedule_name, group_name)

            if not exists:
                note = f"미존재 (그룹: {group_name}, 스케줄: {schedule_name})"
            else:
                note = "존재함"

            return exists, note

        elif resource_type == 'APIGateway':
            exists = check_apigateway_exists(resource_name)
            return exists, "존재함" if exists else "미존재"

        elif resource_type == 'SecretsManager':
            exists = check_secretsmanager_exists(resource_name)
            return exists, "존재함" if exists else "미존재"

        else:
            return False, f"지원하지 않는 리소스 타입: {resource_type}"

    except Exception as e:
        logger.error(f"Error verifying {resource_type} {resource_name}: {e}")
        return False, f"검증 중 오류: {str(e)}"


# ===================================================
# 파일 형식 판별 함수
# ===================================================

def is_csv_file(file_path: str) -> bool:
    """파일이 CSV 형식인지 확인"""
    return file_path.lower().endswith('.csv')


def is_excel_file(file_path: str) -> bool:
    """파일이 Excel 형식인지 확인"""
    return file_path.lower().endswith(('.xlsx', '.xls'))


# ===================================================
# CSV 처리 함수
# ===================================================

def read_csv_checklist(file_path: str) -> List[Dict]:
    """
    CSV 파일에서 마이그레이션 체크리스트 읽기
    """
    logger.info(f"Reading CSV migration checklist from: {file_path}")

    try:
        items = []
        with open(file_path, 'r', encoding='utf-8-sig') as csvfile:
            reader = csv.DictReader(csvfile)

            for row_idx, row in enumerate(reader, start=2):  # start=2 (헤더가 1)
                resource_type = row.get('리소스 타입', '').strip()
                resource_name = row.get('리소스 이름', '').strip()

                # Skip if resource type or name is empty
                if not resource_type or not resource_name:
                    continue

                item = {
                    'row_number': row_idx,
                    'resource_type': resource_type,
                    'resource_name': resource_name,
                    'source_account': row.get('소스 계정', '').strip(),
                    'dest_account': row.get('대상 계정', '').strip(),
                    'current_status': row.get('마이그레이션 상태', '').strip(),
                }
                items.append(item)

        logger.info(f"Read {len(items)} items from CSV checklist")
        return items

    except FileNotFoundError:
        logger.error(f"File not found: {file_path}")
        return []
    except Exception as e:
        logger.error(f"Error reading CSV file: {e}")
        return []


def write_csv_results(file_path: str, results: List[Dict]):
    """
    검증 결과를 CSV 파일로 저장
    """
    logger.info(f"Writing verification results to CSV: {file_path}")

    try:
        # 원본 파일에서 모든 데이터 읽기
        original_data = {}
        input_file = EXCEL_FILE_PATH

        try:
            with open(input_file, 'r', encoding='utf-8-sig') as csvfile:
                reader = csv.DictReader(csvfile)
                fieldnames = reader.fieldnames
                for idx, row in enumerate(reader, start=2):
                    original_data[idx] = row
        except:
            # 원본 파일이 없으면 기본 헤더 사용
            fieldnames = ['리소스 타입', '리소스 이름', '소스 계정', '대상 계정',
                         '마이그레이션 상태', '검증 일시', '비고']

        verification_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        # CSV 파일 작성
        with open(file_path, 'w', newline='', encoding='utf-8-sig') as csvfile:
            writer = csv.DictWriter(csvfile, fieldnames=fieldnames)
            writer.writeheader()

            # 모든 검증 결과를 출력 (원본 데이터 순서 유지)
            for result in results:
                row_num = result['row_number']

                # 원본 데이터가 있으면 사용, 없으면 새로 생성
                if row_num in original_data:
                    row = original_data[row_num]
                else:
                    # 원본에 없는 경우 결과 데이터로 새 행 생성
                    row = {
                        '리소스 타입': result['resource_type'],
                        '리소스 이름': result['resource_name'],
                        '소스 계정': '',
                        '대상 계정': '',
                        '마이그레이션 상태': '',
                        '검증 일시': '',
                        '비고': ''
                    }

                # 검증 결과로 업데이트
                status_icon = '✅' if result['exists'] else '❌'
                row['리소스 타입'] = result['resource_type']
                row['리소스 이름'] = result['resource_name']
                row['마이그레이션 상태'] = f"{status_icon} {result['status']}"
                row['검증 일시'] = verification_time
                row['비고'] = result['notes']

                writer.writerow(row)

        logger.info(f"Successfully saved CSV verification results to: {file_path}")

    except Exception as e:
        logger.error(f"Error writing CSV file: {e}")


# ===================================================
# 엑셀 처리 함수
# ===================================================

def read_excel_checklist(file_path: str) -> List[Dict]:
    """
    엑셀 파일에서 마이그레이션 체크리스트 읽기
    """
    if not OPENPYXL_AVAILABLE:
        logger.error("openpyxl is required for Excel files. Install it with: pip install openpyxl")
        return []

    logger.info(f"Reading Excel migration checklist from: {file_path}")

    try:
        wb = load_workbook(file_path)
        if SHEET_NAME not in wb.sheetnames:
            logger.error(f"Sheet '{SHEET_NAME}' not found in workbook")
            logger.info(f"Available sheets: {wb.sheetnames}")
            return []

        ws = wb[SHEET_NAME]
        items = []

        # Skip header row (row 1)
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            # Skip empty rows
            if not any(row):
                continue

            resource_type = row[COL_RESOURCE_TYPE] if len(row) > COL_RESOURCE_TYPE else None
            resource_name = row[COL_RESOURCE_NAME] if len(row) > COL_RESOURCE_NAME else None

            # Skip if resource type or name is empty
            if not resource_type or not resource_name:
                continue

            item = {
                'row_number': row_idx,
                'resource_type': str(resource_type).strip(),
                'resource_name': str(resource_name).strip(),
                'source_account': str(row[COL_SOURCE_ACCOUNT]).strip() if len(row) > COL_SOURCE_ACCOUNT and row[COL_SOURCE_ACCOUNT] else '',
                'dest_account': str(row[COL_DEST_ACCOUNT]).strip() if len(row) > COL_DEST_ACCOUNT and row[COL_DEST_ACCOUNT] else '',
                'current_status': str(row[COL_STATUS]).strip() if len(row) > COL_STATUS and row[COL_STATUS] else '',
            }
            items.append(item)

        logger.info(f"Read {len(items)} items from Excel checklist")
        return items

    except FileNotFoundError:
        logger.error(f"File not found: {file_path}")
        return []
    except Exception as e:
        logger.error(f"Error reading Excel file: {e}")
        return []


def write_excel_results(file_path: str, results: List[Dict]):
    """
    검증 결과를 엑셀 파일로 저장
    """
    if not OPENPYXL_AVAILABLE:
        logger.error("openpyxl is required for Excel files. Install it with: pip install openpyxl")
        return

    logger.info(f"Writing verification results to Excel: {file_path}")

    try:
        # 원본 파일을 읽어서 수정
        try:
            wb = load_workbook(EXCEL_FILE_PATH)
            ws = wb[SHEET_NAME]
        except:
            # 원본 파일이 없으면 새로 생성
            wb = Workbook()
            ws = wb.active
            ws.title = SHEET_NAME

            # 헤더 작성
            headers = ['리소스 타입', '리소스 이름', '소스 계정', '대상 계정', '마이그레이션 상태', '검증 일시', '비고']
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="CCE5FF", end_color="CCE5FF", fill_type="solid")

        # 결과 업데이트
        verification_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        # 색상 정의
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        for result in results:
            row_num = result['row_number']

            # 리소스 타입과 이름 업데이트 (원본과 동일하게 유지)
            ws.cell(row=row_num, column=COL_RESOURCE_TYPE + 1, value=result['resource_type'])
            ws.cell(row=row_num, column=COL_RESOURCE_NAME + 1, value=result['resource_name'])

            # 상태 업데이트
            status_cell = ws.cell(row=row_num, column=COL_STATUS + 1)
            status_cell.value = result['status']

            # 색상 적용
            if result['exists']:
                status_cell.fill = green_fill
            else:
                status_cell.fill = red_fill

            # 검증 일시 업데이트
            ws.cell(row=row_num, column=COL_VERIFIED_AT + 1, value=verification_time)

            # 비고 업데이트
            ws.cell(row=row_num, column=COL_NOTES + 1, value=result['notes'])

        # 파일 저장
        wb.save(file_path)
        logger.info(f"Successfully saved Excel verification results to: {file_path}")

    except Exception as e:
        logger.error(f"Error writing Excel file: {e}")


# ===================================================
# 통합 처리 함수 (CSV 또는 Excel 자동 감지)
# ===================================================

def read_migration_checklist(file_path: str) -> List[Dict]:
    """
    파일 확장자에 따라 CSV 또는 Excel 파일에서 마이그레이션 체크리스트 읽기
    """
    if is_csv_file(file_path):
        return read_csv_checklist(file_path)
    elif is_excel_file(file_path):
        return read_excel_checklist(file_path)
    else:
        logger.error(f"Unsupported file format: {file_path}. Please use .csv, .xlsx, or .xls")
        return []


def write_verification_results(file_path: str, results: List[Dict]):
    """
    파일 확장자에 따라 CSV 또는 Excel 파일로 검증 결과 저장
    """
    if is_csv_file(file_path):
        write_csv_results(file_path, results)
    elif is_excel_file(file_path):
        write_excel_results(file_path, results)
    else:
        logger.error(f"Unsupported file format: {file_path}. Please use .csv, .xlsx, or .xls")


def print_summary(results: List[Dict]):
    """검증 결과 요약 출력"""
    total = len(results)
    migrated = sum(1 for r in results if r['exists'])
    not_migrated = total - migrated

    print("\n" + "=" * 80)
    print("마이그레이션 검증 결과 요약")
    print("=" * 80)
    print(f"총 리소스 수: {total}")
    print(f"✅ 이관 완료: {migrated} ({migrated/total*100:.1f}%)" if total > 0 else "✅ 이관 완료: 0")
    print(f"❌ 이관 미완료: {not_migrated} ({not_migrated/total*100:.1f}%)" if total > 0 else "❌ 이관 미완료: 0")
    print("=" * 80)

    # 리소스 타입별 통계
    type_stats = {}
    for result in results:
        rtype = result['resource_type']
        if rtype not in type_stats:
            type_stats[rtype] = {'total': 0, 'migrated': 0}
        type_stats[rtype]['total'] += 1
        if result['exists']:
            type_stats[rtype]['migrated'] += 1

    print("\n리소스 타입별 상세:")
    print("-" * 80)
    for rtype, stats in sorted(type_stats.items()):
        total_type = stats['total']
        migrated_type = stats['migrated']
        percent = migrated_type / total_type * 100 if total_type > 0 else 0
        print(f"  {rtype:20s}: {migrated_type:3d}/{total_type:3d} ({percent:5.1f}%)")
    print("-" * 80)

    # 이관 미완료 리소스 목록
    if not_migrated > 0:
        print("\n❌ 이관 미완료 리소스:")
        print("-" * 80)
        for result in results:
            if not result['exists']:
                print(f"  [{result['resource_type']}] {result['resource_name']}")
                if result['notes']:
                    print(f"    → {result['notes']}")
        print("-" * 80)

    print()


# ===================================================
# 메인 함수
# ===================================================

def main():
    """메인 실행 함수"""
    logger.info("=" * 80)
    logger.info("AWS 마이그레이션 검증 스크립트")
    logger.info("=" * 80)
    logger.info(f"대상 계정 프로파일: {DST_PROFILE}")
    logger.info(f"리전: {REGION}")
    logger.info(f"입력 파일: {EXCEL_FILE_PATH}")
    logger.info(f"출력 파일: {OUTPUT_FILE_PATH}")

    # 파일 형식 감지
    if is_csv_file(EXCEL_FILE_PATH):
        logger.info(f"파일 형식: CSV")
    elif is_excel_file(EXCEL_FILE_PATH):
        logger.info(f"파일 형식: Excel")
    else:
        logger.info(f"파일 형식: 알 수 없음")

    logger.info("=" * 80)

    # 1. 파일에서 체크리스트 읽기 (CSV 또는 Excel 자동 감지)
    items = read_migration_checklist(EXCEL_FILE_PATH)
    if not items:
        logger.error("체크리스트가 비어있습니다. 종료합니다.")
        return

    # 2. 각 리소스 검증
    results = []
    logger.info(f"\n리소스 검증 시작 (총 {len(items)}개)")
    logger.info("-" * 80)

    for idx, item in enumerate(items, 1):
        resource_type = item['resource_type']
        resource_name = item['resource_name']

        logger.info(f"[{idx}/{len(items)}] {resource_type}: {resource_name}")

        # 지원하지 않는 리소스 타입 체크
        if resource_type not in SUPPORTED_RESOURCE_TYPES:
            logger.warning(f"  ⚠ 지원하지 않는 리소스 타입: {resource_type}")
            results.append({
                'row_number': item['row_number'],
                'resource_type': resource_type,
                'resource_name': resource_name,
                'exists': False,
                'status': '미지원',
                'notes': f'지원하지 않는 리소스 타입: {resource_type}'
            })
            continue

        # 리소스 존재 여부 확인
        exists, notes = verify_resource(resource_type, resource_name)

        status = '이관 완료' if exists else '이관 미완료'
        status_icon = '✅' if exists else '❌'

        logger.info(f"  {status_icon} {status}")
        if notes:
            logger.info(f"     {notes}")

        results.append({
            'row_number': item['row_number'],
            'resource_type': resource_type,
            'resource_name': resource_name,
            'exists': exists,
            'status': status,
            'notes': notes
        })

    logger.info("-" * 80)

    # 3. 결과를 파일로 저장 (CSV 또는 Excel 자동 감지)
    write_verification_results(OUTPUT_FILE_PATH, results)

    # 4. 요약 출력
    print_summary(results)

    logger.info(f"✅ 검증 완료! 결과 파일: {OUTPUT_FILE_PATH}")


if __name__ == "__main__":
    main()
